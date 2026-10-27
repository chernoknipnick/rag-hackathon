"""Общие фикстуры: фейковый эмбеддер (без скачивания моделей) и набор документов всех форматов."""
import json
import sys
import zlib
from pathlib import Path

import numpy as np
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src import config  # noqa: E402
from src.retrieve.bm25 import tokenize  # noqa: E402


class FakeEmbedder:
    """Хэшированный bag-of-words: детерминированный и достаточный, чтобы проверить пайплайн."""
    dim = 2048

    def encode(self, texts):
        out = np.zeros((len(texts), self.dim), dtype="float32")
        for i, t in enumerate(texts):
            for w in tokenize(t):
                out[i, zlib.crc32(w.encode()) % self.dim] += 1
        return out


def make_docs(root: Path) -> Path:
    import docx
    import pandas as pd

    for c in ("tehnoprom", "logistics", "retail"):
        (root / c).mkdir(parents=True, exist_ok=True)

    d = docx.Document()
    d.add_heading("Договор перевозки с ТрансГруз", level=1)
    d.add_paragraph("Перевозчик обязуется доставлять грузы в согласованные сроки.")
    d.add_heading("5. Ответственность сторон", level=2)
    d.add_paragraph("5.2 За просрочку доставки перевозчик уплачивает штраф 0,1% от стоимости груза за каждый день.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Нарушение", "Штраф"
    t.cell(1, 0).text, t.cell(1, 1).text = "Повреждение груза", "полная стоимость"
    d.add_paragraph("Споры решаются в арбитражном суде.")
    d.save(root / "logistics" / "dogovor_transgruz.docx")

    with pd.ExcelWriter(root / "logistics" / "tarify.xlsx") as w:
        pd.DataFrame([["Тарифная сетка 2024", None, None],
                      ["Маршрут", "Вес, кг", "Цена, руб"],
                      ["Москва — Казань", "500", "12 400"],
                      ["Москва — Тверь", "500", "7 100"]]).to_excel(w, sheet_name="Тарифы", index=False,
                                                                   header=False)

    (root / "retail" / "stock.csv").write_text(
        "Артикул;Название;Остаток\nFC-1;Полетный контроллер;14\nGPS-2;GPS-модуль;3\n", encoding="utf-8")
    (root / "retail" / "cards.json").write_text(json.dumps(
        {"meta": {"source": "shop"}, "catalog": {"items": [
            {"sku": "FC-1", "name": "Полетный контроллер", "price": 5400, "specs": {"weight_g": 30}},
            {"sku": "GPS-2", "name": "GPS-модуль", "price": 2100},
        ]}}, ensure_ascii=False), encoding="utf-8")
    (root / "retail" / "return_policy.md").write_text(
        "# Политика возврата\n\nТовар надлежащего качества можно вернуть в течение 14 дней с момента получения.\n\n"
        "## Возврат после 14 дней\n\nПосле 14 дней возврат возможен только при заводском браке по заключению сервиса.\n",
        encoding="utf-8")
    (root / "retail" / "chat_log.txt").write_text(
        "Клиент: где мой заказ 1234?\nОператор: заказ передан в доставку, срок 3 дня.\n", encoding="utf-8")

    try:
        import fitz
        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), "Technical passport. Cover page.")
        pdf.new_page().insert_text((72, 72), "Stanok TP-400: massa 1200 kg, klass tochnosti P.")
        pdf.save(root / "tehnoprom" / "stanok_tp400.pdf")
    except ImportError:
        (root / "tehnoprom" / "stanok_tp400.md").write_text(
            "# Станок ТП-400\n\nМасса станка ТП-400 1200 кг. Класс точности П.\n\n"
            "## Техническое обслуживание\n\nТО проводится каждые 500 часов работы.\n", encoding="utf-8")
    return root


@pytest.fixture()
def docs(tmp_path, monkeypatch) -> Path:
    root = make_docs(tmp_path / "docs")
    monkeypatch.setattr(config, "DOCS_DIR", root)
    return root


@pytest.fixture()
def index(docs, tmp_path):
    from src.retrieve.index import Index
    idx = Index(embedder=FakeEmbedder(), index_dir=tmp_path / "index")
    idx.ingest_dir(docs)
    return idx
